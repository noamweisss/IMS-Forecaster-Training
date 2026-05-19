import os
import json
from pathlib import Path

# PPTX
try:
    from pptx import Presentation
except ImportError:
    Presentation = None

# PDF
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

# DOCX
try:
    import docx
except ImportError:
    docx = None

def extract_pptx(filepath):
    if not Presentation:
        return {"error": "python-pptx not installed"}
    
    try:
        prs = Presentation(filepath)
        content = []
        for i, slide in enumerate(prs.slides):
            slide_content = {"slide_number": i + 1, "text": [], "notes": ""}
            
            # Extract text from shapes
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_content["text"].append(shape.text.strip())
            
            # Extract notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_content["notes"] = notes
            
            content.append(slide_content)
        return {"type": "pptx", "slides": content}
    except Exception as e:
        return {"type": "pptx", "error": str(e)}

def extract_pdf(filepath):
    if not fitz:
        return {"error": "pymupdf not installed"}
    
    try:
        doc = fitz.open(filepath)
        content = []
        for i, page in enumerate(doc):
            text = page.get_text().strip()
            if text:
                content.append({"page_number": i + 1, "text": text})
        return {"type": "pdf", "pages": content}
    except Exception as e:
        return {"type": "pdf", "error": str(e)}

def extract_docx(filepath):
    if not docx:
        return {"error": "python-docx not installed"}
    
    try:
        doc = docx.Document(filepath)
        content = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                content.append(text)
        return {"type": "docx", "paragraphs": content}
    except Exception as e:
        return {"type": "docx", "error": str(e)}

def main():
    source_dir = Path("source_files")
    output_dir = Path("module_output")
    output_dir.mkdir(exist_ok=True)
    
    results = {}
    
    if not source_dir.exists():
        print(f"Source directory {source_dir} not found.")
        return
        
    for filepath in source_dir.iterdir():
        if filepath.is_file():
            print(f"Processing: {filepath.name}")
            ext = filepath.suffix.lower()
            if ext == ".pptx":
                results[filepath.name] = extract_pptx(filepath)
            elif ext == ".pdf":
                results[filepath.name] = extract_pdf(filepath)
            elif ext == ".docx":
                results[filepath.name] = extract_docx(filepath)
            else:
                print(f"Skipping {filepath.name} (unsupported format)")
                
    output_file = output_dir / "_extraction.json"
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
        
    print(f"Extraction complete. Saved to {output_file}")

if __name__ == "__main__":
    main()
