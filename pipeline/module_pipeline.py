#!/usr/bin/env python3
"""
Course Factory — Module Pipeline
IMS Aviation Forecaster Certification (and other courses)

Converts a folder of PPTX / PDF / DOCX files into one Moodle-ready module:
  - Visual-forward HTML lesson pages (with inline SVG diagrams)
  - Moodle XML quiz files (per lesson + combined module quiz)
  - Module overview page
  - Import instructions (README)

Usage:
    python module_pipeline.py \
        --input   ./aviation_module_1 \
        --name    "Meteorological Observations for Aviation" \
        --number  1 \
        --course  "Aviation Weather Forecasting Certification" \
        --audience "Professional meteorologists seeking Aviation Forecaster certification (IMS)" \
        --output  ./aviation_module_1_output

Requirements:
    pip install anthropic python-pptx pymupdf python-docx

Environment:
    ANTHROPIC_API_KEY must be set.
"""

import argparse
import asyncio
import json
import os
import re
import zipfile
from pathlib import Path

import anthropic

# ── Optional library imports ────────────────────────────────────────────────────

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False
    print("⚠  python-pptx not found — PPTX files will be skipped.")

try:
    import fitz  # PyMuPDF
    HAS_PDF = True
except ImportError:
    HAS_PDF = False
    print("⚠  PyMuPDF not found — PDF files will be skipped.")

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    print("⚠  python-docx not found — DOCX files will be skipped.")

client = anthropic.Anthropic()
MODEL  = "claude-sonnet-4-20250514"


# ══════════════════════════════════════════════════════════════════════════════
#  STAGE 1: EXTRACTION
# ══════════════════════════════════════════════════════════════════════════════

def extract_pptx(filepath: Path) -> dict:
    if not HAS_PPTX:
        return {"filename": filepath.name, "format": "pptx", "error": "python-pptx not installed", "slides": []}

    prs = Presentation(filepath)
    slides = []

    for i, slide in enumerate(prs.slides):
        sd = {
            "slide_number": i + 1,
            "title":        "",
            "body":         [],
            "notes":        "",
            "has_images":   False,
            "has_charts":   False,
        }

        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                sd["has_images"] = True
            if hasattr(shape, "chart"):
                sd["has_charts"] = True
            if not shape.has_text_frame:
                continue

            raw_text = shape.text_frame.text.strip()
            if not raw_text:
                continue

            is_title = (
                shape.name == "Title"
                or (
                    hasattr(shape, "placeholder_format")
                    and shape.placeholder_format
                    and shape.placeholder_format.idx == 0
                )
            )

            if is_title:
                sd["title"] = raw_text
            else:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        sd["body"].append({"level": para.level, "text": t})

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes and "Click to edit" not in notes:
                sd["notes"] = notes

        slides.append(sd)

    return {
        "filename":    filepath.name,
        "format":      "pptx",
        "slide_count": len(slides),
        "slides":      slides,
    }


def extract_pdf(filepath: Path) -> dict:
    if not HAS_PDF:
        return {"filename": filepath.name, "format": "pdf", "error": "PyMuPDF not installed", "pages": []}

    doc   = fitz.open(str(filepath))
    pages = []

    for page_num, page in enumerate(doc):
        blocks = page.get_text("blocks")
        images = page.get_images()

        title     = ""
        body_list = []

        for block in blocks:
            if len(block) >= 5:
                txt = block[4].strip()
                if txt:
                    if not title:
                        title = txt[:120]
                    else:
                        body_list.append(txt[:600])

        pages.append({
            "page_number": page_num + 1,
            "title":       title,
            "body":        body_list,
            "has_images":  len(images) > 0,
        })

    doc.close()

    return {
        "filename":   filepath.name,
        "format":     "pdf",
        "page_count": len(pages),
        "pages":      pages,
    }


def extract_docx(filepath: Path) -> dict:
    if not HAS_DOCX:
        return {"filename": filepath.name, "format": "docx", "error": "python-docx not installed", "sections": []}

    doc      = DocxDocument(filepath)
    sections = []
    current  = {"heading": filepath.stem, "paragraphs": [], "has_tables": False}

    for para in doc.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue
        if para.style.name.startswith("Heading"):
            if current["paragraphs"]:
                sections.append(current)
            current = {"heading": txt, "paragraphs": [], "has_tables": False}
        else:
            current["paragraphs"].append(txt)

    for table in doc.tables:
        current["has_tables"] = True

    if current["paragraphs"] or current["heading"] != filepath.stem:
        sections.append(current)

    return {
        "filename":      filepath.name,
        "format":        "docx",
        "section_count": len(sections),
        "sections":      sections,
    }


def extract_file(filepath: Path) -> dict:
    ext = filepath.suffix.lower()
    if ext == ".pptx":
        return extract_pptx(filepath)
    elif ext == ".pdf":
        return extract_pdf(filepath)
    elif ext in (".docx", ".doc"):
        return extract_docx(filepath)
    else:
        return {"filename": filepath.name, "error": f"Unsupported format: {ext}"}


def extract_all(input_dir: Path) -> list[dict]:
    supported = {".pptx", ".pdf", ".docx", ".doc"}
    files = sorted(f for f in input_dir.iterdir() if f.suffix.lower() in supported)

    print(f"\n  Found {len(files)} source files:")
    for f in files:
        print(f"    {f.suffix.upper()[1:]:5}  {f.name}")

    results = []
    for f in files:
        print(f"  Extracting {f.name}...", end=" ", flush=True)
        results.append(extract_file(f))
        print("✓")

    return results


# ══════════════════════════════════════════════════════════════════════════════
#  STAGE 2: BLUEPRINT (lesson structure within the pre-determined module)
# ══════════════════════════════════════════════════════════════════════════════

BLUEPRINT_SYSTEM = """\
You are a senior instructional designer creating certification-level courses for technical experts at a national meteorological service.

CONTEXT:
- Courses are self-paced and standalone. No instructor. No class. No supplementary material exists.
- Learners are experienced professional meteorologists — not students. Do not condescend.
- Each module is an independent unit reviewed and approved separately by institutional stakeholders.
- Visual learning is a first-class concern: every abstract concept should have a diagram, table, or annotated figure.

YOUR TASK:
Given extracted content from several source files that together form ONE pre-determined module,
design the lesson structure for that module.

GUIDELINES:
- 3–6 lessons per module is the right range for a 90-minute module.
- Lesson objectives must be verb-led, specific, and measurable (Bloom's taxonomy action verbs).
- For visual_elements_needed: be specific. Not "a diagram" — "a process flowchart showing the sequence of METAR observation and encoding steps" or "a comparison table of ICAO vs WMO cloud classification codes."
- Flag any topic that requires real photographic or instrument imagery that must be sourced externally.

OUTPUT: Return ONLY valid JSON, no markdown fences, no commentary:
{
  "module_title": "...",
  "module_description": "2–3 sentences for the learner. What this module covers and why it matters operationally.",
  "estimated_duration_minutes": 90,
  "competencies": [
    "Interpret METAR observations in degraded visibility conditions",
    "..."
  ],
  "lessons": [
    {
      "lesson_id": "L1",
      "lesson_title": "...",
      "learning_objectives": [
        "Decode all mandatory METAR groups according to ICAO Annex 3",
        "..."
      ],
      "source_files": ["filename.pptx", "other.pdf"],
      "key_concepts": ["METAR", "SPECI", "TAF"],
      "visual_elements_needed": [
        "Annotated METAR string showing each group with color-coded labels",
        "Flowchart: when to issue SPECI vs routine METAR"
      ],
      "external_imagery_needed": [
        "Photograph of convective cloud development over an airport"
      ],
      "estimated_duration_minutes": 20
    }
  ],
  "module_quiz": {
    "description": "End-of-module summative assessment",
    "question_count": 15,
    "passing_score_percent": 75
  }
}
"""


def generate_blueprint(extracted: list[dict], module_name: str, module_number: int,
                        course_context: str, audience: str) -> dict:

    # Build a compact summary (not the full dump — too large for the blueprint pass)
    summaries = []
    for doc in extracted:
        s = {"filename": doc["filename"], "format": doc.get("format", "?")}

        if doc.get("format") == "pptx":
            s["slide_count"]  = doc.get("slide_count", 0)
            s["slide_titles"] = [sl["title"] for sl in doc.get("slides", []) if sl.get("title")]
            # First 4 slides with body bullets for content sampling
            s["sample"] = [
                {"title": sl["title"],
                 "bullets": [b["text"] for b in sl.get("body", [])[:4]]}
                for sl in doc.get("slides", [])[:4]
            ]

        elif doc.get("format") == "pdf":
            s["page_count"] = doc.get("page_count", 0)
            s["titles"]     = [p["title"][:80] for p in doc.get("pages", []) if p.get("title")]

        elif doc.get("format") == "docx":
            s["section_count"] = doc.get("section_count", 0)
            s["headings"]      = [sec["heading"] for sec in doc.get("sections", []) if sec.get("heading")]

        if doc.get("error"):
            s["error"] = doc["error"]

        summaries.append(s)

    prompt = (
        f"Course: {course_context}\n"
        f"Module {module_number}: {module_name}\n"
        f"Audience: {audience}\n\n"
        f"Source files ({len(extracted)} total):\n"
        f"{json.dumps(summaries, indent=2, ensure_ascii=False)}\n\n"
        f"Design the lesson structure. Return JSON only."
    )

    print("  Calling Claude for blueprint...", end=" ", flush=True)
    resp = client.messages.create(
        model=MODEL,
        max_tokens=2500,
        system=BLUEPRINT_SYSTEM,
        messages=[{"role": "user", "content": prompt}],
    )
    text = _strip_fences(resp.content[0].text)
    result = json.loads(text)
    print("✓")
    return result


# ══════════════════════════════════════════════════════════════════════════════
#  STAGE 3: CONTENT GENERATION (one lesson at a time, run in parallel)
# ══════════════════════════════════════════════════════════════════════════════

LESSON_SYSTEM = """\
You are a senior instructional designer and technical writer. You create professional certification content for expert practitioners — not introductory material.

AUDIENCE: Experienced professional meteorologists at a national meteorological service. They are technically literate and time-pressured. Respect their expertise.

TONE: Authoritative, precise, direct. No filler sentences. No "in this lesson we will learn..." type openers. Start with content.

VISUAL-FIRST APPROACH:
Every abstract concept must have a visual anchor. You will generate:
  - SVG diagrams for process flows, decision trees, timelines, system diagrams
  - HTML comparison tables for reference data
  - Styled callout boxes (note, warning, key formula) for important items
  - A visual element is required every 200–300 words of prose

SVG RULES (critical — these SVGs embed directly in HTML and will be rendered
inside Moodle, which strips our CSS variables. NEVER reference var(--anything)
inside SVG attributes — only literal hex colors):
  - viewBox="0 0 680 H" where H fits the content
  - Flat colors only (no gradients)
  - Font: sans-serif, 14px labels, 12px subtitles
  - Include role="img" with <title> and <desc>
  - ALL fill, stroke, and color attributes MUST be literal hex (#RRGGBB).
    Never `fill="var(--bg)"`, never `stroke="var(--text)"`. The page's CSS
    custom properties do not exist inside Moodle pages.
  - Color palette (use these literals):
      accent blue   #185FA5      text dark      #1A1917
      teal          #0F6E56      muted gray     #6B6A65
      amber         #BA7517      border gray    #E2DFD6
      coral         #993C1D      surface cream  #F6F5F0
      key purple    #534AB7
  - Light background fills (for boxes inside SVGs): #E6F1FB, #E1F5EE,
    #FAEEDA, #EEEDFE. Never use pure white or black as a fill.
  - All text inside SVGs uses dark hex (#1A1917 or similar) so it reads on
    any page background.

HTML CONTENT RULES:
  - Return body content only. NO <html>, <head>, <body>, <!DOCTYPE>, or
    <style> tags. The pipeline wraps your output in the surrounding
    document and applies the design system CSS automatically (every rule
    is scoped to the `.ims-lesson` wrapper).
  - Use semantic elements: <section>, <figure>, <figcaption>, <table>
  - Prose paragraphs: 2–4 sentences max
  - Use the callout classes: <div class="callout callout-note">, callout-warning, callout-key
  - Provide a <div class="lesson-summary"> at the end with 4–6 key takeaway bullets
  - Where a real photograph or instrument image is needed (not generatable as SVG),
    insert: <div class="image-needed" data-description="DESCRIPTION OF NEEDED IMAGE">
            <span>📷 Image needed: DESCRIPTION</span></div>
    This flags it for human sourcing during review.

QUIZ QUESTIONS:
  - Write at application level (Bloom's), not recall. Scenarios preferred over definitions.
  - All 4 options must be plausible. Wrong answers should reflect real misconceptions.
  - Include a clear rationale explaining WHY the correct answer is right.

OUTPUT: Return ONLY valid JSON (no markdown fences):
{
  "lesson_title": "...",
  "html_content": "<section>full lesson HTML here</section>",
  "quiz_questions": [
    {
      "type": "multiple_choice",
      "question": "...",
      "options": {"A": "...", "B": "...", "C": "...", "D": "..."},
      "correct": "B",
      "rationale": "...",
      "difficulty": "medium",
      "objective_id": "L1"
    }
  ]
}
"""


def _gather_source_content(lesson: dict, extracted: list[dict], max_chars: int = 7000) -> str:
    """Pull relevant text from source files for this lesson's content generation."""
    source_files = lesson.get("source_files", [])
    chunks = []

    for doc in extracted:
        # Include if no filter set, or this file is relevant to the lesson
        is_relevant = not source_files or any(sf in doc["filename"] for sf in source_files)
        if not is_relevant:
            continue

        fmt = doc.get("format", "?")

        if fmt == "pptx":
            lines = []
            for sl in doc.get("slides", []):
                line = f"\nSlide {sl['slide_number']}: {sl.get('title','(no title)')}"
                for b in sl.get("body", []):
                    line += f"\n{'  ' * b['level']}• {b['text']}"
                if sl.get("notes"):
                    line += f"\n  [Notes: {sl['notes'][:300]}]"
                lines.append(line)
            chunks.append(f"=== {doc['filename']} ({doc.get('slide_count',0)} slides) ===\n" + "\n".join(lines[:30]))

        elif fmt == "pdf":
            lines = []
            for pg in doc.get("pages", [])[:20]:
                body = " ".join(pg.get("body", [])[:6])
                if body.strip():
                    lines.append(f"\nPage {pg['page_number']}: {pg.get('title','')} — {body[:400]}")
            chunks.append(f"=== {doc['filename']} ({doc.get('page_count',0)} pages) ===\n" + "\n".join(lines))

        elif fmt == "docx":
            lines = []
            for sec in doc.get("sections", [])[:15]:
                lines.append(f"\n## {sec.get('heading','')}")
                lines.append("\n".join(sec.get("paragraphs", [])[:4]))
            chunks.append(f"=== {doc['filename']} ===\n" + "\n".join(lines))

    combined = "\n\n".join(chunks)
    return combined[:max_chars]


# Light-mode hex values for every CSS custom property declared in
# config/design_system.css. Used to scrub `var(--xxx)` references out of
# SVG attributes before the HTML is uploaded to Moodle, which does not
# preserve our CSS variable scope (see docs/journal.md 2026-05-19).
CSS_VAR_HEX_FALLBACKS = {
    "--bg":            "#FFFFFF",
    "--surface":       "#F6F5F0",
    "--border":        "#E2DFD6",
    "--text":          "#1A1917",
    "--muted":         "#6B6A65",
    "--accent":        "#185FA5",
    "--accent-light":  "#E6F1FB",
    "--note-bg":       "#E1F5EE",
    "--note-border":   "#0F6E56",
    "--warn-bg":       "#FAEEDA",
    "--warn-border":   "#BA7517",
    "--key-bg":        "#EEEDFE",
    "--key-border":    "#534AB7",
}


def _inline_css_vars_in_svgs(html: str) -> str:
    """Replace `var(--name)` references inside <svg>...</svg> blocks with
    literal hex from CSS_VAR_HEX_FALLBACKS. Leaves var() references outside
    SVGs untouched — those live in real CSS rules and resolve normally."""

    var_pattern = re.compile(r"var\(\s*(--[A-Za-z0-9_-]+)\s*(?:,[^)]*)?\)")

    def scrub(match: re.Match) -> str:
        svg_block = match.group(0)
        def replace_var(v: re.Match) -> str:
            name = v.group(1)
            return CSS_VAR_HEX_FALLBACKS.get(name, v.group(0))
        return var_pattern.sub(replace_var, svg_block)

    return re.sub(r"<svg\b.*?</svg>", scrub, html, flags=re.DOTALL | re.IGNORECASE)


async def generate_lesson_content(lesson: dict, extracted: list[dict],
                                   blueprint: dict, course_context: str,
                                   audience: str) -> dict:

    source_text = _gather_source_content(lesson, extracted)
    n_questions = max(4, min(8, lesson.get("estimated_duration_minutes", 20) // 4))

    prompt = (
        f"Course: {course_context}\n"
        f"Module: {blueprint.get('module_title','')}\n"
        f"Audience: {audience}\n\n"
        f"LESSON TO CREATE:\n"
        f"  ID:              {lesson['lesson_id']}\n"
        f"  Title:           {lesson['lesson_title']}\n"
        f"  Duration:        {lesson.get('estimated_duration_minutes', 20)} minutes\n"
        f"  Objectives:\n"
        + "\n".join(f"    - {o}" for o in lesson.get("learning_objectives", []))
        + f"\n  Key concepts:    {', '.join(lesson.get('key_concepts', []))}\n"
        f"  Diagrams needed:\n"
        + "\n".join(f"    - {v}" for v in lesson.get("visual_elements_needed", []))
        + f"\n  External imagery needed:\n"
        + "\n".join(f"    - {e}" for e in lesson.get("external_imagery_needed", []))
        + f"\n\nSOURCE MATERIAL:\n{source_text}\n\n"
        f"Generate the full lesson and {n_questions} quiz questions. Return JSON only."
    )

    loop = asyncio.get_event_loop()
    resp = await loop.run_in_executor(
        None,
        lambda: client.messages.create(
            model=MODEL,
            max_tokens=5000,
            system=LESSON_SYSTEM,
            messages=[{"role": "user", "content": prompt}],
        ),
    )

    text = _strip_fences(resp.content[0].text)

    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        # Salvage: find outermost JSON object
        m = re.search(r'\{.*\}', text, re.DOTALL)
        if m:
            try:
                parsed = json.loads(m.group())
            except Exception:
                return {
                    "lesson_title": lesson["lesson_title"],
                    "error":        "JSON parse failed",
                    "raw_response": text[:800],
                }
        else:
            return {
                "lesson_title": lesson["lesson_title"],
                "error":        "JSON parse failed",
                "raw_response": text[:800],
            }

    # Post-process: scrub any var(--xxx) the model slipped into SVG attributes.
    if isinstance(parsed.get("html_content"), str):
        parsed["html_content"] = _inline_css_vars_in_svgs(parsed["html_content"])

    return parsed


async def generate_all_lessons(blueprint: dict, extracted: list[dict],
                                course_context: str, audience: str) -> list[dict]:
    lessons = blueprint.get("lessons", [])
    print(f"\n  Generating content for {len(lessons)} lessons in parallel...")

    tasks = [
        generate_lesson_content(l, extracted, blueprint, course_context, audience)
        for l in lessons
    ]

    results = await asyncio.gather(*tasks, return_exceptions=True)

    # Unwrap exceptions into error dicts
    cleaned = []
    for r, l in zip(results, lessons):
        if isinstance(r, Exception):
            cleaned.append({"lesson_title": l["lesson_title"], "error": str(r)})
        else:
            cleaned.append(r)
        status = "✗ ERROR" if isinstance(r, Exception) else "✓"
        print(f"    {status}  {l['lesson_title']}")

    return cleaned


# ══════════════════════════════════════════════════════════════════════════════
#  STAGE 4: ASSEMBLY
# ══════════════════════════════════════════════════════════════════════════════

def _get_page_css() -> str:
    css_path = Path(__file__).parent.parent / "config" / "design_system.css"
    if css_path.exists():
        with open(css_path, "r", encoding="utf-8") as f:
            return f.read()
    return ""


def _html_page(title: str, breadcrumb: str, content: str) -> str:
    """Standalone preview HTML. Body content is wrapped in
    `<div class="ims-lesson">` so the CSS (scoped to that class) applies."""
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<style>
body {{ margin: 0; background: #fff; }}
@media (prefers-color-scheme: dark) {{ body {{ background: #1a1917; }} }}
{_get_page_css()}
</style>
</head>
<body>
<div class="ims-lesson">
<p class="breadcrumb">{breadcrumb}</p>
<h1>{title}</h1>
{content}
</div>
</body>
</html>"""


def _quiz_xml(questions: list[dict], category_name: str) -> str:
    parts = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        "<quiz>",
        f"  <question type=\"category\"><category><text>$course$/top/{category_name}</text></category></question>",
    ]

    for i, q in enumerate(questions, 1):
        if q.get("type") != "multiple_choice":
            continue
        opts    = q.get("options", {})
        correct = q.get("correct", "A")

        parts.append(f"""  <question type="multichoice">
    <name><text>Q{i:02d} — {q.get('question','')[:60]}</text></name>
    <questiontext format="html">
      <text><![CDATA[<p>{q.get('question','')}</p>]]></text>
    </questiontext>
    <generalfeedback format="html">
      <text><![CDATA[<p><strong>Rationale:</strong> {q.get('rationale','')}</p>]]></text>
    </generalfeedback>
    <defaultgrade>1</defaultgrade>
    <single>true</single>
    <shuffleanswers>true</shuffleanswers>""")

        for key, text in opts.items():
            frac = "100" if key == correct else "0"
            parts.append(f"""    <answer fraction="{frac}" format="html">
      <text><![CDATA[<p>{text}</p>]]></text>
      <feedback format="html"><text></text></feedback>
    </answer>""")

        parts.append("  </question>")

    parts.append("</quiz>")
    return "\n".join(parts)


def assemble_overview(blueprint: dict, module_number: int) -> str:
    lessons      = blueprint.get("lessons", [])
    competencies = blueprint.get("competencies", [])

    comp_html    = "".join(f"<li>{c}</li>" for c in competencies)
    quiz_info    = blueprint.get("module_quiz", {})

    lessons_html = ""
    for i, l in enumerate(lessons, 1):
        objs = "".join(f"<li>{o}</li>" for o in l.get("learning_objectives", []))
        lessons_html += f"""
<div style="border:1px solid var(--border); border-radius:8px; padding:1rem 1.25rem; margin:0.75rem 0;">
  <div style="display:flex; justify-content:space-between; align-items:baseline;">
    <strong>Lesson {i}: {l.get('lesson_title','')}</strong>
    <span style="color:var(--muted); font-size:0.85rem;">~{l.get('estimated_duration_minutes',20)} min</span>
  </div>
  <ul style="margin:0.5rem 0 0; padding-left:1.4rem; font-size:0.9rem;">{objs}</ul>
</div>"""

    body = f"""
<p class="meta">
  Estimated duration: ~{blueprint.get('estimated_duration_minutes', 90)} min &nbsp;|&nbsp;
  {len(lessons)} lessons &nbsp;|&nbsp;
  Module quiz: {quiz_info.get('question_count', 15)} questions,
  pass at {quiz_info.get('passing_score_percent', 75)}%
</p>

<p>{blueprint.get('module_description', '')}</p>

<div class="objectives">
  <strong>By completing this module you will be certified to:</strong>
  <ul>{comp_html}</ul>
</div>

<h2>Lessons</h2>
{lessons_html}
"""

    return _html_page(
        blueprint.get("module_title", f"Module {module_number}"),
        f"Module {module_number} overview",
        body,
    )


def _safe_slug(title: str) -> str:
    slug = re.sub(r"[^\w\s-]", "", title)
    slug = re.sub(r"\s+", "_", slug.strip()).lower()
    return slug[:60]


# ══════════════════════════════════════════════════════════════════════════════
#  MAIN ORCHESTRATOR
# ══════════════════════════════════════════════════════════════════════════════

async def run_pipeline(args: argparse.Namespace) -> None:
    input_dir  = Path(args.input)
    output_dir = Path(args.output)

    # Create output tree
    lessons_dir = output_dir / "lessons"
    lessons_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "assets" / "images").mkdir(parents=True, exist_ok=True)

    _header("Course Factory — Module Pipeline")
    print(f"  Course:  {args.course}")
    print(f"  Module:  {args.number}  —  {args.name}")
    print(f"  Output:  {output_dir.resolve()}")

    # ── Stage 1: Extract ──────────────────────────────────────────────────────
    _header("Stage 1 / Extraction")
    extracted = extract_all(input_dir)
    _save_json(output_dir / "_extraction.json", extracted)
    print(f"  Saved raw extraction → _extraction.json")

    # ── Stage 2: Blueprint ────────────────────────────────────────────────────
    _header("Stage 2 / Blueprint")
    blueprint = generate_blueprint(
        extracted, args.name, args.number, args.course, args.audience
    )
    _save_json(output_dir / "module_structure.json", blueprint)
    print(f"  {len(blueprint.get('lessons', []))} lessons defined → module_structure.json")

    # ── Stage 3: Content generation ───────────────────────────────────────────
    _header("Stage 3 / Content generation")
    lesson_contents = await generate_all_lessons(
        blueprint, extracted, args.course, args.audience
    )

    # ── Stage 4: Assembly ─────────────────────────────────────────────────────
    _header("Stage 4 / Assembly")

    # Module overview
    overview_html = assemble_overview(blueprint, args.number)
    _write(output_dir / "00_module_overview.html", overview_html)
    print("  00_module_overview.html")

    all_questions  = []
    image_requests = []  # Collect image-needed markers for the summary

    for i, (lesson_def, content) in enumerate(
        zip(blueprint["lessons"], lesson_contents), 1
    ):
        if content.get("error"):
            print(f"  ✗  Lesson {i} — error: {content['error'][:80]}")
            continue

        lesson_id = lesson_def.get("lesson_id", f"L{i}")
        slug      = _safe_slug(lesson_def["lesson_title"])

        # Lesson HTML
        objectives_html = "".join(
            f"<li>{o}</li>" for o in lesson_def.get("learning_objectives", [])
        )
        body = f"""
<div class="objectives">
  <strong>By the end of this lesson you will be able to:</strong>
  <ul>{objectives_html}</ul>
</div>
{content.get('html_content', '<p>Content generation failed.</p>')}
"""
        lesson_page = _html_page(
            content.get("lesson_title", lesson_def["lesson_title"]),
            f"Module {args.number} / Lesson {i}",
            body,
        )
        lesson_filename = f"{i:02d}_{slug}.html"
        _write(lessons_dir / lesson_filename, lesson_page)

        # Quiz XML
        questions = content.get("quiz_questions", [])
        if questions:
            xml_str  = _quiz_xml(questions, f"Module{args.number}_{lesson_id}")
            quiz_file = f"{i:02d}_{slug}_quiz.xml"
            _write(lessons_dir / quiz_file, xml_str)
            all_questions.extend(questions)
            print(f"  lessons/{lesson_filename}  +  {quiz_file} ({len(questions)} Qs)")
        else:
            print(f"  lessons/{lesson_filename}")

        # Collect image-needed annotations from the HTML
        for desc in re.findall(r'data-description="([^"]+)"', content.get("html_content", "")):
            image_requests.append(f"Lesson {i} ({lesson_def['lesson_title']}): {desc}")

    # Combined module quiz
    if all_questions:
        combined_xml = _quiz_xml(all_questions, f"Module{args.number}_Combined")
        _write(output_dir / "module_quiz_all_questions.xml", combined_xml)
        print(f"\n  module_quiz_all_questions.xml  ({len(all_questions)} total questions)")

    # Image sourcing list
    if image_requests:
        image_doc = "# Images needed for human sourcing\n\n"
        image_doc += "The following visuals could not be auto-generated and must be sourced from the IMS media library or photographed:\n\n"
        for req in image_requests:
            image_doc += f"- {req}\n"
        _write(output_dir / "IMAGES_TO_SOURCE.md", image_doc)
        print(f"\n  IMAGES_TO_SOURCE.md  ({len(image_requests)} items flagged)")

    # README
    _write(output_dir / "README.md", _build_readme(
        blueprint, args, all_questions, image_requests
    ))

    # ZIP the whole output
    zip_path = output_dir.parent / f"{output_dir.name}.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for fp in output_dir.rglob("*"):
            if fp.is_file():
                zf.write(fp, fp.relative_to(output_dir.parent))

    _header("Done")
    print(f"  Output folder:  {output_dir.resolve()}")
    print(f"  ZIP package:    {zip_path.resolve()}")
    print()


def _build_readme(blueprint, args, all_questions, image_requests) -> str:
    module_title = blueprint.get("module_title", args.name)
    lessons      = blueprint.get("lessons", [])
    quiz_info    = blueprint.get("module_quiz", {})

    lines = [
        f"# {module_title}",
        f"## Module {args.number} — Generated Content Package",
        "",
        "### Files in this package",
        "",
        "| File | What to do with it in Moodle |",
        "|------|------------------------------|",
        "| `00_module_overview.html` | Add as a **Page** activity at the top of the module section |",
    ]

    for i, l in enumerate(lessons, 1):
        slug = _safe_slug(l["lesson_title"])
        lines.append(f"| `lessons/{i:02d}_{slug}.html` | Add as a **Page** activity (Lesson {i}) |")
        lines.append(f"| `lessons/{i:02d}_{slug}_quiz.xml` | Import via Question Bank → Import → Moodle XML |")

    lines += [
        "| `module_quiz_all_questions.xml` | All questions combined — create a Quiz activity and pull from this bank |",
        "| `module_structure.json` | Course outline — keep for reference / pipeline re-runs |",
        "",
        "### How to import into Moodle (step by step)",
        "1. In your course, turn editing on and add a new **Topic section** for this module.",
        "2. Add a **Page** activity → paste the content of `00_module_overview.html` (body only) into the HTML source editor.",
        "3. Repeat for each `lessons/NN_*.html` file.",
        "4. Go to **Question Bank** → **Import** → choose **Moodle XML format** → upload `module_quiz_all_questions.xml`.",
        "5. Add a **Quiz** activity → under 'Edit quiz', add questions from the imported bank.",
        "6. Set the passing grade to "
        f"{quiz_info.get('passing_score_percent', 75)}% in the quiz settings.",
        "",
        "### Module summary",
        f"- **Lessons:** {len(lessons)}",
        f"- **Estimated duration:** ~{blueprint.get('estimated_duration_minutes', 90)} min",
        f"- **Quiz questions:** {len(all_questions)}",
        f"- **Passing score:** {quiz_info.get('passing_score_percent', 75)}%",
        f"- **Images to source:** {len(image_requests)} (see `IMAGES_TO_SOURCE.md`)",
        "",
        "### Generation parameters",
        f"- Course: {args.course}",
        f"- Audience: {args.audience}",
        f"- Model: {MODEL}",
    ]

    return "\n".join(lines)


# ══════════════════════════════════════════════════════════════════════════════
#  UTILITIES
# ══════════════════════════════════════════════════════════════════════════════

def _strip_fences(text: str) -> str:
    text = re.sub(r"^```[a-z]*\s*", "", text.strip())
    text = re.sub(r"\s*```$", "", text)
    return text.strip()


def _save_json(path: Path, data) -> None:
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def _write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)


def _header(text: str) -> None:
    print(f"\n{'─'*55}")
    print(f"  {text}")
    print(f"{'─'*55}")


# ══════════════════════════════════════════════════════════════════════════════
#  ENTRY POINT
# ══════════════════════════════════════════════════════════════════════════════

def main():
    p = argparse.ArgumentParser(
        description="Course Factory — Module Pipeline",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    p.add_argument("--input",    required=True, metavar="DIR",    help="Folder with source files (PPTX/PDF/DOCX)")
    p.add_argument("--name",     required=True, metavar="NAME",   help="Pre-determined module name")
    p.add_argument("--number",   required=True, type=int,         help="Module number within the course (e.g. 1)")
    p.add_argument("--course",   required=True, metavar="COURSE", help="Course name / context")
    p.add_argument("--audience", required=True, metavar="TEXT",   help="Target audience description")
    p.add_argument("--output",   required=True, metavar="DIR",    help="Output folder (will be created)")
    args = p.parse_args()

    asyncio.run(run_pipeline(args))


if __name__ == "__main__":
    main()
