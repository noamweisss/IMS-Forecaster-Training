#!/usr/bin/env python3
"""
extract_sources.py — Read PPTX/PDF/DOCX source files and emit a single
JSON dump containing every slide, page, and section's text content, plus
any embedded images written out to disk.

Pure Python. No LLM. No API key needed. Run this before handing the
content over to a Claude Code agent for lesson generation — the agent
reads the JSON instead of opening binary files itself.

Usage:
    python pipeline/extract_sources.py \
        --input  courses/aviation-weather/source_files/module-2 \
        --output courses/aviation-weather/module-2/_extraction.json

By default, embedded images are extracted into
`<output-parent>/extracted_images/<source-stem>/` and indexed in the JSON
under each slide/page's `images` array. Pass `--no-images` to skip image
extraction, or `--images-dir` to write them somewhere else.
"""

import argparse
import hashlib
import io
import json
import os
import re
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

try:
    import fitz  # PyMuPDF
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from PIL import Image
    HAS_PILLOW = True
except ImportError:
    HAS_PILLOW = False


def _safe_stem(name: str) -> str:
    """Sanitize a filename stem for use as a directory name."""
    stem = re.sub(r"[^a-zA-Z0-9._-]+", "_", name).strip("_")
    return stem or "unnamed"


def _image_dimensions(data: bytes) -> tuple[int | None, int | None]:
    if not HAS_PILLOW or not data:
        return None, None
    try:
        with Image.open(io.BytesIO(data)) as img:
            return img.width, img.height
    except Exception:
        return None, None


def _write_image(
    data: bytes,
    ext: str,
    images_root: Path,
    source_stem: str,
    name_stem: str,
    img_idx: int,
    hash_cache: dict,
    module_dir: Path,
) -> dict | None:
    """Write image bytes under `images_root/<source_stem>/` and return a
    metadata dict. Dedupes by md5 within the same source — repeated
    bytes (e.g. a logo on every slide) reuse the first on-disk file.

    Returns None for empty input.
    """
    if not data:
        return None

    digest = hashlib.md5(data).hexdigest()[:8]
    if digest in hash_cache:
        return hash_cache[digest]

    ext = (ext or "bin").lstrip(".").lower()
    if ext == "jpg":
        ext = "jpeg"

    out_dir = images_root / source_stem
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{name_stem}_img{img_idx}.{ext}"
    out_path = out_dir / filename
    out_path.write_bytes(data)

    width, height = _image_dimensions(data)
    rel_path = os.path.relpath(out_path, start=module_dir).replace(os.sep, "/")

    meta = {
        "path":   rel_path,
        "format": ext,
        "width":  width,
        "height": height,
        "hash":   digest,
    }
    hash_cache[digest] = meta
    return meta


def extract_pptx(
    filepath: Path,
    images_root: Path | None = None,
    module_dir: Path | None = None,
) -> dict:
    if not HAS_PPTX:
        return {"filename": filepath.name, "format": "pptx",
                "error": "python-pptx not installed", "slides": []}

    prs    = Presentation(filepath)
    slides = []
    source_stem = _safe_stem(filepath.stem)
    hash_cache: dict = {}

    for i, slide in enumerate(prs.slides):
        sd = {
            "slide_number": i + 1,
            "title":        "",
            "body":         [],
            "notes":        "",
            "has_images":   False,
            "has_charts":   False,
            "images":       [],
        }

        img_idx = 0
        for shape in slide.shapes:
            shape_img = getattr(shape, "image", None)
            is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE or shape_img is not None
            if is_picture:
                sd["has_images"] = True
                if images_root is not None and shape_img is not None:
                    try:
                        img_idx += 1
                        meta = _write_image(
                            shape_img.blob,
                            shape_img.ext,
                            images_root,
                            source_stem,
                            f"slide{i + 1:02d}",
                            img_idx,
                            hash_cache,
                            module_dir,
                        )
                        if meta:
                            sd["images"].append(meta)
                    except Exception as exc:
                        print(f"    warn: could not extract image on slide {i + 1} "
                              f"of {filepath.name}: {exc}")
            if getattr(shape, "has_chart", False):
                sd["has_charts"] = True
            if not shape.has_text_frame:
                continue

            raw_text = shape.text_frame.text.strip()
            if not raw_text:
                continue

            is_title = (
                shape.name == "Title"
                or (
                    getattr(shape, "is_placeholder", False)
                    and shape.placeholder_format.idx == 0
                )
            )

            if is_title:
                sd["title"] = raw_text
            else:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        sd["body"].append({"level": para.level, "text": t})

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes and "Click to edit" not in notes:
                sd["notes"] = notes

        slides.append(sd)

    return {
        "filename":    filepath.name,
        "format":      "pptx",
        "slide_count": len(slides),
        "slides":      slides,
    }


def extract_pdf(
    filepath: Path,
    images_root: Path | None = None,
    module_dir: Path | None = None,
) -> dict:
    if not HAS_PDF:
        return {"filename": filepath.name, "format": "pdf",
                "error": "PyMuPDF not installed", "pages": []}

    doc   = fitz.open(str(filepath))
    pages = []
    source_stem = _safe_stem(filepath.stem)
    hash_cache: dict = {}

    for page_num, page in enumerate(doc):
        blocks = page.get_text("blocks")
        raw_images = page.get_images(full=True)

        title     = ""
        body_list = []

        for block in blocks:
            if len(block) >= 5:
                txt = block[4].strip()
                if txt:
                    if not title:
                        title = txt[:120]
                    else:
                        body_list.append(txt[:600])

        page_dict = {
            "page_number": page_num + 1,
            "title":       title,
            "body":        body_list,
            "has_images":  len(raw_images) > 0,
            "images":      [],
        }

        if images_root is not None:
            for img_idx, img_info in enumerate(raw_images, start=1):
                xref = img_info[0]
                try:
                    extracted = doc.extract_image(xref)
                except Exception as exc:
                    print(f"    warn: could not extract image {img_idx} on page "
                          f"{page_num + 1} of {filepath.name}: {exc}")
                    continue
                meta = _write_image(
                    extracted.get("image", b""),
                    extracted.get("ext", "png"),
                    images_root,
                    source_stem,
                    f"page{page_num + 1:02d}",
                    img_idx,
                    hash_cache,
                    module_dir,
                )
                if meta:
                    page_dict["images"].append(meta)

        pages.append(page_dict)

    doc.close()

    return {
        "filename":   filepath.name,
        "format":     "pdf",
        "page_count": len(pages),
        "pages":      pages,
    }


def extract_docx(
    filepath: Path,
    images_root: Path | None = None,
    module_dir: Path | None = None,
) -> dict:
    if not HAS_DOCX:
        return {"filename": filepath.name, "format": "docx",
                "error": "python-docx not installed", "sections": []}

    doc      = DocxDocument(filepath)
    sections = []
    current  = {"heading": filepath.stem, "paragraphs": [], "has_tables": False}

    for para in doc.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue
        if para.style.name.startswith("Heading"):
            if current["paragraphs"]:
                sections.append(current)
            current = {"heading": txt, "paragraphs": [], "has_tables": False}
        else:
            current["paragraphs"].append(txt)

    for _ in doc.tables:
        current["has_tables"] = True

    if current["paragraphs"] or current["heading"] != filepath.stem:
        sections.append(current)

    # DOCX images are attached at the document level rather than to a
    # specific section — attributing an inline_shape back to its host
    # paragraph requires walking the XML, which isn't worth the
    # complexity for the few DOCX sources this pipeline sees. The agent
    # can still browse `module_images` and choose the relevant figure.
    module_images: list = []
    if images_root is not None:
        source_stem = _safe_stem(filepath.stem)
        hash_cache: dict = {}
        img_idx = 0
        for _, part in doc.part.related_parts.items():
            content_type = getattr(part, "content_type", "") or ""
            if not content_type.startswith("image/"):
                continue
            blob = getattr(part, "blob", None)
            if not blob:
                continue
            img_idx += 1
            ext = content_type.split("/", 1)[1].split("+", 1)[0]
            try:
                meta = _write_image(
                    blob,
                    ext,
                    images_root,
                    source_stem,
                    "doc",
                    img_idx,
                    hash_cache,
                    module_dir,
                )
            except Exception as exc:
                print(f"    warn: could not extract image {img_idx} from "
                      f"{filepath.name}: {exc}")
                continue
            if meta:
                module_images.append(meta)

    return {
        "filename":      filepath.name,
        "format":        "docx",
        "section_count": len(sections),
        "sections":      sections,
        "has_images":    bool(module_images),
        "module_images": module_images,
    }


def extract_file(
    filepath: Path,
    images_root: Path | None = None,
    module_dir: Path | None = None,
) -> dict:
    ext = filepath.suffix.lower()
    if ext == ".pptx":
        return extract_pptx(filepath, images_root, module_dir)
    if ext == ".pdf":
        return extract_pdf(filepath, images_root, module_dir)
    if ext in (".docx", ".doc"):
        return extract_docx(filepath, images_root, module_dir)
    return {"filename": filepath.name, "error": f"Unsupported format: {ext}"}


def extract_all(
    input_dir: Path,
    images_root: Path | None = None,
    module_dir: Path | None = None,
) -> list[dict]:
    supported = {".pptx", ".pdf", ".docx", ".doc"}
    files = sorted(f for f in input_dir.iterdir()
                   if f.is_file() and f.suffix.lower() in supported)

    print(f"Found {len(files)} source file(s) in {input_dir}:")
    for f in files:
        print(f"  {f.suffix.upper()[1:]:5}  {f.name}")

    results = []
    for f in files:
        print(f"  extracting {f.name} ...", end=" ", flush=True)
        results.append(extract_file(f, images_root, module_dir))
        print("done")

    return results


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--input", required=True,
                   help="Folder containing the source PPTX/PDF/DOCX files.")
    p.add_argument("--output", required=True,
                   help="Output JSON path (typically <module>/_extraction.json).")
    p.add_argument("--no-images", action="store_true",
                   help="Skip extracting embedded images to disk.")
    p.add_argument("--images-dir", default=None,
                   help="Where to write extracted images. Defaults to "
                        "<output-parent>/extracted_images. Ignored if "
                        "--no-images is set.")
    args = p.parse_args()

    in_dir = Path(args.input)
    if not in_dir.is_dir():
        raise SystemExit(f"Input folder not found: {in_dir}")

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    module_dir = out.parent

    if args.no_images:
        images_root = None
    else:
        images_root = Path(args.images_dir) if args.images_dir else module_dir / "extracted_images"
        images_root.mkdir(parents=True, exist_ok=True)

    extracted = extract_all(in_dir, images_root, module_dir)
    with out.open("w", encoding="utf-8") as f:
        json.dump(extracted, f, indent=2, ensure_ascii=False)

    msg = f"\nWrote {out}  ({len(extracted)} file(s) extracted"
    if images_root is not None:
        msg += f", images under {images_root.relative_to(module_dir) if images_root.is_relative_to(module_dir) else images_root}"
    msg += ")"
    print(msg)


if __name__ == "__main__":
    main()
